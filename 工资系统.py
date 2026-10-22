import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
import pandas as pd
from datetime import datetime
import os
import re

class 工资系统:
    def __init__(self, root):
        self.root = root
        self.root.title("冲压车间工资录入系统")
        self.root.geometry("1050x780")
        self.root.resizable(True, True)
        
        # 数据变量
        self.产品库 = {}  # {产品名: 单价}
        self.员工列表 = []
        self.当前员工 = tk.StringVar()
        self.当前日期 = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        self.搜索关键词 = tk.StringVar()
        self.搜索关键词.trace('w', self.刷新产品列表)
        self.选中产品 = tk.StringVar()
        
        # 文件路径
        self.源文件 = "示例数据.xlsx"  # 实际使用时替换为真实计数表文件名
        self.输出文件 = "工资记录.xlsx"
        self.员工配置文件 = "员工名单.txt"
        self.产品配置文件 = "产品库.txt"
        
        # 加载数据
        self.加载员工列表()
        self.加载产品库()
        
        # 创建界面
        self.创建界面()
        
    def 加载员工列表(self):
        """加载员工列表：优先从txt文件读，如果没有则从Excel读"""
        if os.path.exists(self.员工配置文件):
            try:
                with open(self.员工配置文件, 'r', encoding='utf-8') as f:
                    self.员工列表 = [line.strip() for line in f.readlines() if line.strip()]
                if self.员工列表:
                    self.当前员工.set(self.员工列表[0])
                    return
            except:
                pass
        
        # 从Excel读取员工
        try:
            if os.path.exists(self.源文件):
                xl = pd.ExcelFile(self.源文件)
                self.员工列表 = [s for s in xl.sheet_names if s not in ["Sheet17", "Sheet18"]]
                if self.员工列表:
                    self.当前员工.set(self.员工列表[0])
                    self.保存员工列表()
                    return
        except:
            pass
        
        # 默认员工列表（示例数据，实际使用时请通过界面添加）
        self.员工列表 = [
            "员工A", "员工B", "员工C", "员工D", "员工E",
            "员工F", "员工G", "员工H", "员工I", "员工J",
            "员工K", "员工L", "员工M", "员工N", "员工O", "员工P"
        ]
        if self.员工列表:
            self.当前员工.set(self.员工列表[0])
        self.保存员工列表()
    
    def 保存员工列表(self):
        """保存员工列表到txt文件"""
        try:
            with open(self.员工配置文件, 'w', encoding='utf-8') as f:
                for 员工 in self.员工列表:
                    f.write(员工 + '\n')
        except:
            pass
    
    def 加载产品库(self):
        """加载产品库：优先从txt文件读，如果没有则从Excel读"""
        if os.path.exists(self.产品配置文件):
            try:
                with open(self.产品配置文件, 'r', encoding='utf-8') as f:
                    self.产品库 = {}
                    for line in f.readlines():
                        if line.strip():
                            parts = line.strip().split('|||')
                            if len(parts) == 2:
                                self.产品库[parts[0]] = float(parts[1])
                if self.产品库:
                    return
            except:
                pass
        
        # 从Excel读取
        try:
            if not os.path.exists(self.源文件):
                messagebox.showwarning("警告", f"找不到 {self.源文件}，请确保文件在程序同目录下")
                return
            
            df = pd.read_excel(self.源文件, sheet_name=None, header=None)
            self.产品库 = {}
            
            for sheet_name, sheet_data in df.items():
                if sheet_name in ["Sheet17", "Sheet18"]:
                    continue
                if sheet_data.empty:
                    continue
                for idx, row in sheet_data.iterrows():
                    if idx < 2:
                        continue
                    try:
                        产品名 = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
                        单价 = float(row.iloc[3]) if pd.notna(row.iloc[3]) else 0
                        if 产品名 and 单价 > 0 and 产品名 != "nan":
                            if 产品名 not in self.产品库:
                                self.产品库[产品名] = 单价
                    except:
                        continue
            
            self.保存产品库()
            messagebox.showinfo("加载完成", f"已加载 {len(self.产品库)} 种产品")
        except Exception as e:
            messagebox.showerror("错误", f"加载产品库失败：{str(e)}")
    
    def 保存产品库(self):
        """保存产品库到txt文件"""
        try:
            with open(self.产品配置文件, 'w', encoding='utf-8') as f:
                for 产品, 单价 in sorted(self.产品库.items()):
                    f.write(f"{产品}|||{单价}\n")
        except:
            pass
    
    def 创建界面(self):
        """创建主界面（大字体版）"""
        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 定义大字体
        large_font = ("微软雅黑", 13, "bold")
        normal_font = ("微软雅黑", 12)
        list_font = ("微软雅黑", 13)
        
        # ========== 顶部：员工选择 + 管理 ==========
        top_frame = ttk.LabelFrame(main_frame, text="👤 员工管理", padding="15")
        top_frame.pack(fill=tk.X, pady=(0, 15))
        
        # 第一行：选择员工
        row1 = ttk.Frame(top_frame)
        row1.pack(fill=tk.X, pady=(0, 10))
        
        ttk.Label(row1, text="当前员工：", font=large_font).pack(side=tk.LEFT)
        self.员工下拉 = ttk.Combobox(row1, textvariable=self.当前员工, 
                                     values=self.员工列表, width=18, 
                                     font=normal_font, state="readonly")
        self.员工下拉.pack(side=tk.LEFT, padx=(10, 30))
        
        ttk.Label(row1, text="日期：", font=large_font).pack(side=tk.LEFT)
        日期输入 = ttk.Entry(row1, textvariable=self.当前日期, width=14, font=normal_font)
        日期输入.pack(side=tk.LEFT, padx=(10, 10))
        ttk.Label(row1, text="（格式：2026-08-09）", font=normal_font).pack(side=tk.LEFT)
        
        # 第二行：员工管理按钮
        row2 = ttk.Frame(top_frame)
        row2.pack(fill=tk.X, pady=(5, 0))
        
        ttk.Button(row2, text="➕ 添加员工", command=self.添加员工, width=14).pack(side=tk.LEFT, padx=(0, 10))
        ttk.Button(row2, text="➖ 删除员工", command=self.删除员工, width=14).pack(side=tk.LEFT, padx=(0, 10))
        ttk.Button(row2, text="📋 查看所有员工", command=self.查看所有员工, width=16).pack(side=tk.LEFT, padx=(0, 20))
        
        ttk.Label(row2, text=f"共 {len(self.员工列表)} 人", font=large_font).pack(side=tk.LEFT)
        
        # ========== 中部：产品搜索与选择 + 管理 ==========
        middle_frame = ttk.LabelFrame(main_frame, text="🔍 产品管理", padding="15")
        middle_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 15))
        
        # 搜索行
        search_frame = ttk.Frame(middle_frame)
        search_frame.pack(fill=tk.X, pady=(0, 12))
        
        ttk.Label(search_frame, text="搜索产品：", font=large_font).pack(side=tk.LEFT)
        搜索框 = ttk.Entry(search_frame, textvariable=self.搜索关键词, width=35, font=normal_font)
        搜索框.pack(side=tk.LEFT, padx=(10, 15))
        
        ttk.Label(search_frame, text=f"共 {len(self.产品库)} 种产品", font=large_font).pack(side=tk.LEFT, padx=(10, 0))
        
        # 产品管理按钮
        btn_frame = ttk.Frame(search_frame)
        btn_frame.pack(side=tk.RIGHT)
        
        ttk.Button(btn_frame, text="➕ 添加产品", command=self.添加产品, width=12).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(btn_frame, text="➖ 删除产品", command=self.删除产品, width=12).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(btn_frame, text="🔄 刷新", command=self.刷新产品库, width=10).pack(side=tk.LEFT)
        
        # 产品列表（加大字体）
        list_frame = ttk.Frame(middle_frame)
        list_frame.pack(fill=tk.BOTH, expand=True)
        
        scrollbar = ttk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.产品列表 = tk.Listbox(list_frame, yscrollcommand=scrollbar.set, 
                                   font=list_font, selectmode=tk.SINGLE,
                                   height=12)
        self.产品列表.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=self.产品列表.yview)
        self.产品列表.bind('<<ListboxSelect>>', self.产品选中)
        self.产品列表.bind('<Double-Button-1>', self.双击确认)
        
        self.刷新产品列表()
        
        # ========== 底部：录入信息 ==========
        bottom_frame = ttk.LabelFrame(main_frame, text="📦 录入信息", padding="15")
        bottom_frame.pack(fill=tk.X)
        
        product_frame = ttk.Frame(bottom_frame)
        product_frame.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(product_frame, text="已选产品：", font=large_font).pack(side=tk.LEFT)
        self.产品标签 = ttk.Label(product_frame, text="（请从列表中选择）", 
                                 foreground="blue", font=large_font)
        self.产品标签.pack(side=tk.LEFT, padx=(10, 0))
        
        price_frame = ttk.Frame(bottom_frame)
        price_frame.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(price_frame, text="单　　价：", font=large_font).pack(side=tk.LEFT)
        self.单价标签 = ttk.Label(price_frame, text="--", foreground="green", 
                                  font=("微软雅黑", 15, "bold"))
        self.单价标签.pack(side=tk.LEFT, padx=(10, 30))
        
        ttk.Label(price_frame, text="数　　量：", font=large_font).pack(side=tk.LEFT)
        self.数量输入 = ttk.Entry(price_frame, width=14, font=normal_font)
        self.数量输入.pack(side=tk.LEFT, padx=(10, 10))
        ttk.Label(price_frame, text="件", font=large_font).pack(side=tk.LEFT)
        
        # 操作按钮（加大）
        btn_row = ttk.Frame(bottom_frame)
        btn_row.pack(fill=tk.X, pady=(10, 0))
        
        ttk.Button(btn_row, text="✅ 录入工资", command=self.录入工资, width=16).pack(side=tk.LEFT, padx=(0, 15))
        ttk.Button(btn_row, text="📊 生成汇总表", command=self.生成汇总表, width=16).pack(side=tk.LEFT, padx=(0, 15))
        ttk.Button(btn_row, text="📈 查询工资", command=self.查询工资, width=16).pack(side=tk.LEFT)
        
        # 状态栏（加大字体）
        self.状态栏 = ttk.Label(self.root, text="就绪", relief=tk.SUNKEN, 
                                anchor=tk.W, padding=(10, 5), font=normal_font)
        self.状态栏.pack(side=tk.BOTTOM, fill=tk.X)
    
    # ==================== 员工管理功能 ====================
    
    def 添加员工(self):
        """添加新员工"""
        姓名 = simpledialog.askstring("添加员工", "请输入新员工姓名：")
        if not 姓名 or not 姓名.strip():
            return
        姓名 = 姓名.strip()
        
        if 姓名 in self.员工列表:
            messagebox.showwarning("提示", f"员工「{姓名}」已存在！")
            return
        
        self.员工列表.append(姓名)
        self.保存员工列表()
        self.刷新员工下拉()
        self.当前员工.set(姓名)
        self.状态栏.config(text=f"✅ 已添加员工：{姓名}")
        messagebox.showinfo("成功", f"已添加员工「{姓名}」")
    
    def 删除员工(self):
        """删除当前选中的员工"""
        if len(self.员工列表) <= 1:
            messagebox.showwarning("提示", "至少保留一名员工！")
            return
        
        员工 = self.当前员工.get()
        if not 员工:
            return
        
        if not messagebox.askyesno("确认删除", f"确定要删除员工「{员工}」吗？\n（工资记录不会自动删除）"):
            return
        
        self.员工列表.remove(员工)
        self.保存员工列表()
        self.刷新员工下拉()
        self.当前员工.set(self.员工列表[0])
        self.状态栏.config(text=f"已删除员工：{员工}")
    
    def 查看所有员工(self):
        """弹出窗口显示所有员工"""
        if not self.员工列表:
            messagebox.showinfo("员工列表", "暂无员工")
            return
        内容 = "📋 员工名单（共 " + str(len(self.员工列表)) + " 人）\n" + "="*30 + "\n"
        for i, 员工 in enumerate(self.员工列表, 1):
            内容 += f"{i}. {员工}\n"
        messagebox.showinfo("员工列表", 内容)
    
    def 刷新员工下拉(self):
        """刷新员工下拉框"""
        self.员工下拉['values'] = self.员工列表
    
    # ==================== 产品管理功能 ====================
    
    def 添加产品(self):
        """添加新产品"""
        dialog = tk.Toplevel(self.root)
        dialog.title("添加产品")
        dialog.geometry("550x250")
        dialog.resizable(False, False)
        dialog.transient(self.root)
        dialog.grab_set()
        
        # 加大字体
        large_font = ("微软雅黑", 13, "bold")
        normal_font = ("微软雅黑", 12)
        
        ttk.Label(dialog, text="产品名称：", font=large_font).grid(row=0, column=0, padx=15, pady=20, sticky=tk.W)
        产品输入 = ttk.Entry(dialog, width=50, font=normal_font)
        产品输入.grid(row=0, column=1, padx=15, pady=20)
        产品输入.focus_set()
        
        ttk.Label(dialog, text="产品单价：", font=large_font).grid(row=1, column=0, padx=15, pady=15, sticky=tk.W)
        单价输入 = ttk.Entry(dialog, width=15, font=normal_font)
        单价输入.grid(row=1, column=1, padx=15, pady=15, sticky=tk.W)
        ttk.Label(dialog, text="元/件", font=large_font).grid(row=1, column=2, padx=(0, 15), pady=15, sticky=tk.W)
        
        def 确认添加():
            产品名 = 产品输入.get().strip()
            单价_str = 单价输入.get().strip()
            
            if not 产品名:
                messagebox.showwarning("提示", "请输入产品名称")
                return
            if not 单价_str:
                messagebox.showwarning("提示", "请输入产品单价")
                return
            
            try:
                单价 = float(单价_str)
                if 单价 <= 0:
                    messagebox.showwarning("提示", "单价必须大于0")
                    return
            except:
                messagebox.showwarning("提示", "请输入有效的数字")
                return
            
            if 产品名 in self.产品库:
                if not messagebox.askyesno("提示", f"产品「{产品名}」已存在，是否覆盖单价？"):
                    return
            
            self.产品库[产品名] = 单价
            self.保存产品库()
            self.刷新产品列表()
            self.状态栏.config(text=f"✅ 已添加产品：{产品名}（单价：{单价}）")
            dialog.destroy()
            messagebox.showinfo("成功", f"已添加产品「{产品名}」\n单价：{单价} 元/件")
        
        btn_frame = ttk.Frame(dialog)
        btn_frame.grid(row=2, column=0, columnspan=3, pady=25)
        ttk.Button(btn_frame, text="✅ 确认添加", command=确认添加, width=14).pack(side=tk.LEFT, padx=10)
        ttk.Button(btn_frame, text="取消", command=dialog.destroy, width=10).pack(side=tk.LEFT, padx=10)
    
    def 删除产品(self):
        """删除当前选中的产品"""
        产品 = self.选中产品.get()
        if not 产品 or 产品 not in self.产品库:
            messagebox.showwarning("提示", "请先从列表中选中要删除的产品")
            return
        
        if not messagebox.askyesno("确认删除", f"确定要删除产品「{产品}」吗？\n单价：{self.产品库[产品]} 元/件"):
            return
        
        del self.产品库[产品]
        self.保存产品库()
        self.选中产品.set("")
        self.产品标签.config(text="（请从列表中选择）", foreground="blue")
        self.单价标签.config(text="--", foreground="green")
        self.刷新产品列表()
        self.状态栏.config(text=f"已删除产品：{产品}")
        messagebox.showinfo("成功", f"已删除产品「{产品}」")
    
    # ==================== 产品列表刷新 ====================
    
    def 刷新产品列表(self, *args):
        self.产品列表.delete(0, tk.END)
        关键词 = self.搜索关键词.get().strip()
        
        if 关键词 == "":
            count = 0
            for 产品 in sorted(self.产品库.keys())[:500]:
                self.产品列表.insert(tk.END, 产品)
                count += 1
            if len(self.产品库) > 500:
                self.产品列表.insert(tk.END, f"... 共 {len(self.产品库)} 种产品，请输入关键词筛选")
        else:
            matches = []
            for 产品 in self.产品库.keys():
                if 关键词.lower() in 产品.lower():
                    matches.append(产品)
                    if len(matches) >= 200:
                        break
            for 产品 in sorted(matches):
                self.产品列表.insert(tk.END, 产品)
            if not matches:
                self.产品列表.insert(tk.END, f"（没有找到包含「{关键词}」的产品）")
    
    def 产品选中(self, event):
        selection = self.产品列表.curselection()
        if selection:
            产品 = self.产品列表.get(selection[0])
            if 产品 in self.产品库:
                self.选中产品.set(产品)
                self.产品标签.config(text=产品, foreground="blue")
                self.单价标签.config(text=f"{self.产品库[产品]:.4f} 元/件", foreground="green")
    
    def 双击确认(self, event):
        self.数量输入.focus_set()
    
    # ==================== 核心功能：录入工资 ====================
    
    def 录入工资(self):
        员工 = self.当前员工.get()
        日期 = self.当前日期.get()
        产品 = self.选中产品.get()
        数量 = self.数量输入.get()
        
        if not 员工:
            messagebox.showwarning("提示", "请选择员工")
            return
        
        if not 产品 or 产品 not in self.产品库:
            messagebox.showwarning("提示", "请从列表中选择一个产品")
            return
        
        try:
            数量_val = float(数量)
            if 数量_val <= 0:
                messagebox.showwarning("提示", "数量必须大于0")
                return
        except:
            messagebox.showwarning("提示", "请输入有效的数量")
            return
        
        单价 = self.产品库[产品]
        金额 = 数量_val * 单价
        
        if not messagebox.askyesno("确认录入", 
            f"确认录入以下信息？\n\n"
            f"👤 员工：{员工}\n"
            f"📅 日期：{日期}\n"
            f"📦 产品：{产品}\n"
            f"🔢 数量：{数量_val}\n"
            f"💰 单价：{单价:.4f}\n"
            f"💵 工钱：{金额:.2f} 元"):
            return
        
        try:
            self.保存记录(员工, 日期, 产品, 数量_val, 单价, 金额)
            self.状态栏.config(text=f"✅ 已录入：{员工} | {产品} | {数量_val}件 | {金额:.2f}元")
            self.数量输入.delete(0, tk.END)
            messagebox.showinfo("成功", f"✅ 录入成功！\n工钱：{金额:.2f} 元")
        except Exception as e:
            messagebox.showerror("错误", f"保存失败：{str(e)}")
    
    # ==================== 核心功能：保存记录 ====================
    
    def 保存记录(self, 员工, 日期, 产品, 数量, 单价, 金额):
        """保存记录到Excel（两行标题）"""
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        
        # ---- 读取已有数据 ----
        if os.path.exists(self.输出文件):
            try:
                df = pd.read_excel(self.输出文件, sheet_name=员工, engine='openpyxl', 
                                   skiprows=2, header=None)
                df.columns = [0, 1, 2, 3, 4]
            except:
                df = pd.DataFrame(columns=[0, 1, 2, 3, 4])
        else:
            df = pd.DataFrame(columns=[0, 1, 2, 3, 4])
        
        # 转换日期
        try:
            日期_obj = datetime.strptime(日期, "%Y-%m-%d")
            日期数字 = 日期_obj.strftime("%Y-%m-%d")
        except:
            日期数字 = 日期
        
        # ---- 只添加新的一行 ----
        new_row = pd.DataFrame({
            0: [日期数字],
            1: [产品],
            2: [数量],
            3: [单价],
            4: [金额]
        })
        
        # ---- 自动调整列宽函数 ----
        def 自动调整列宽(worksheet):
            for col in worksheet.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column].width = adjusted_width
        
        # ---- 写入Excel ----
        try:
            if os.path.exists(self.输出文件):
                existing_sheets = []
                try:
                    with pd.ExcelFile(self.输出文件, engine='openpyxl') as xl:
                        existing_sheets = xl.sheet_names
                except:
                    pass
                
                if 员工 in existing_sheets:
                    # 工作表已存在，只追加新的一行
                    wb = load_workbook(self.输出文件)
                    ws = wb[员工]
                    ws.append([日期数字, 产品, 数量, 单价, 金额])
                    自动调整列宽(ws)
                    wb.save(self.输出文件)
                    wb.close()
                    return
                else:
                    df = pd.concat([df, new_row], ignore_index=True)
                    with pd.ExcelWriter(self.输出文件, engine='openpyxl', mode='a', 
                                        if_sheet_exists='replace') as writer:
                        df.to_excel(writer, sheet_name=员工, index=False, 
                                   startrow=2, header=False)
                        workbook = writer.book
                        worksheet = workbook[员工]
                        worksheet['A1'] = 员工
                        worksheet['B1'] = "冲压车间"
                        worksheet['A2'] = "日期"
                        worksheet['B2'] = "名称"
                        worksheet['C2'] = "数量"
                        worksheet['D2'] = "单价"
                        worksheet['E2'] = "金额"
                        自动调整列宽(worksheet)
                        workbook.save(self.输出文件)
            else:
                df = pd.concat([df, new_row], ignore_index=True)
                with pd.ExcelWriter(self.输出文件, engine='openpyxl') as writer:
                    df.to_excel(writer, sheet_name=员工, index=False, 
                               startrow=2, header=False)
                    workbook = writer.book
                    worksheet = workbook[员工]
                    worksheet['A1'] = 员工
                    worksheet['B1'] = "冲压车间"
                    worksheet['A2'] = "日期"
                    worksheet['B2'] = "名称"
                    worksheet['C2'] = "数量"
                    worksheet['D2'] = "单价"
                    worksheet['E2'] = "金额"
                    自动调整列宽(worksheet)
                    workbook.save(self.输出文件)
                    
        except Exception as e:
            try:
                if os.path.exists(self.输出文件):
                    with pd.ExcelWriter(self.输出文件, engine='openpyxl', mode='a', 
                                        if_sheet_exists='replace') as writer:
                        df.to_excel(writer, sheet_name=员工, index=False, header=False)
                else:
                    with pd.ExcelWriter(self.输出文件, engine='openpyxl') as writer:
                        df.to_excel(writer, sheet_name=员工, index=False, header=False)
            except:
                messagebox.showerror("错误", f"保存失败：{str(e)}")
                raise
    
    # ==================== 核心功能：汇总和查询 ====================
    
    def 生成汇总表(self):
        if not os.path.exists(self.输出文件):
            messagebox.showwarning("提示", "还没有录入数据，请先录入")
            return
        
        try:
            xl = pd.ExcelFile(self.输出文件)
            汇总数据 = []
            
            for sheet in xl.sheet_names:
                try:
                    df = pd.read_excel(self.输出文件, sheet_name=sheet, engine='openpyxl', 
                                       skiprows=2, header=None)
                    总数量 = df[2].sum()
                    总金额 = df[4].sum()
                    汇总数据.append({
                        "员工": sheet,
                        "总数量": int(总数量),
                        "总工资": round(总金额, 2)
                    })
                except:
                    continue
            
            if not 汇总数据:
                messagebox.showwarning("提示", "没有数据")
                return
            
            汇总df = pd.DataFrame(汇总数据)
            汇总df = 汇总df.sort_values("总工资", ascending=False)
            total_row = pd.DataFrame({
                "员工": ["【合计】"],
                "总数量": [汇总df["总数量"].sum()],
                "总工资": [round(汇总df["总工资"].sum(), 2)]
            })
            汇总df = pd.concat([汇总df, total_row], ignore_index=True)
            
            汇总文件 = f"工资汇总_{datetime.now().strftime('%Y%m%d')}.xlsx"
            with pd.ExcelWriter(汇总文件, engine='openpyxl') as writer:
                汇总df.to_excel(writer, sheet_name="汇总", index=False)
            
            messagebox.showinfo("完成", f"✅ 汇总表已生成！\n文件：{汇总文件}")
            self.状态栏.config(text=f"已生成汇总表：{汇总文件}")
        except Exception as e:
            messagebox.showerror("错误", f"生成汇总失败：{str(e)}")
    
    def 查询工资(self):
        if not os.path.exists(self.输出文件):
            messagebox.showwarning("提示", "还没有录入数据")
            return
        
        员工 = simpledialog.askstring("查询工资", "请输入员工姓名（留空查全部）：")
        try:
            xl = pd.ExcelFile(self.输出文件)
            result = "📊 工资查询结果\n" + "="*40 + "\n"
            total = 0
            
            for sheet in xl.sheet_names:
                if 员工 and sheet != 员工:
                    continue
                try:
                    df = pd.read_excel(self.输出文件, sheet_name=sheet, engine='openpyxl', 
                                       skiprows=2, header=None)
                    金额 = df[4].sum()
                    if 金额 > 0:
                        result += f"{sheet}：{round(金额, 2)} 元\n"
                        total += 金额
                except:
                    continue
            
            if 员工:
                result += "="*40 + f"\n合计：{round(total, 2)} 元"
            else:
                result += "="*40 + f"\n全厂合计：{round(total, 2)} 元"
            
            messagebox.showinfo("查询结果", result)
        except Exception as e:
            messagebox.showerror("错误", str(e))
    
    def 刷新产品库(self):
        self.产品库.clear()
        self.加载产品库()
        self.刷新产品列表()
        self.状态栏.config(text=f"已刷新，共 {len(self.产品库)} 种产品")


if __name__ == "__main__":
    root = tk.Tk()
    app = 工资系统(root)
    root.mainloop()